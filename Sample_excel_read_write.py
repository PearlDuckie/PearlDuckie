from openpyxl import Workbook, load_workbook
import pandas as pd
import csv

# Writing to an Excel file
def write_excel(filename):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws['A1'] = "Name"
    ws['B1'] = "Age"
    ws.append(["Alice", 24])
    ws.append(["Bob", 30])
    wb.save(filename)
    print(f"Data written to {filename}")

# Reading from an Excel file
def read_excel(filename):
    wb = load_workbook(filename)
    ws = wb.active
    for row in ws.iter_rows(values_only=True):
        print(row)

# Reading columns from a CSV file with header in row 15 using pandas
def read_csv_columns_pandas(filename, header_row=15):
    # header=14 because pandas uses 0-based index (row 15 is index 14)
    df = pd.read_csv(filename, header=header_row-1)
    for col in df.columns:
        print(f"Column: {col}")
        print(df[col].tolist())

# Alternative: Reading columns using csv module
def read_csv_columns_csv_module(filename, header_row=15):
    with open(filename, mode='r', newline='', encoding='utf-8') as f:
        reader = csv.reader(f)
        # Skip first header_row - 1 rows
        for _ in range(header_row - 1):
            next(reader, None)
        headers = next(reader, None)
        if headers is None:
            print("CSV does not contain enough rows for a header.")
            return
        columns = {header: [] for header in headers}
        for row in reader:
            for i, header in enumerate(headers):
                if i < len(row):
                    columns[header].append(row[i])
                else:
                    columns[header].append('')
        for col, values in columns.items():
            print(f"Column: {col}")
            print(values)

if __name__ == "__main__":
    file = "sample.xlsx"
    csv_file = "sample.csv"  # Replace with your CSV filename

    #write_excel(file)

    print("Reading data from the Excel file:")
    read_excel(file)

    print("\nReading data as columns from the CSV file (header at row 15):")
    # Choose one of the following:
    read_csv_columns_pandas(csv_file, header_row=15)
    # Or, if you prefer the csv module:
    # read_csv_columns_csv_module(csv_file, header_row=15)
